"""Geração do relatório em Excel: o que a equipe jurídica realmente abre."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .estado import CONCLUIDOS

CAMPOS = [
    ("tribunal", "Tribunal"),
    ("grau", "Grau"),
    ("classe", "Classe"),
    ("orgao_julgador", "Órgão julgador"),
    ("data_ajuizamento", "Ajuizamento"),
    ("assuntos", "Assuntos"),
    ("qtd_movimentos", "Qtd. movimentos"),
    ("ultima_movimentacao", "Última movimentação"),
    ("data_ultima_movimentacao", "Data última movimentação"),
    ("documento", "Documento baixado"),
]


def _tabela(linhas: list[dict]) -> pd.DataFrame:
    registros = []
    for linha in linhas:
        reg = {"Processo": linha["numero"], "Status": linha["status"]}
        for chave, titulo in CAMPOS:
            reg[titulo] = linha["dados"].get(chave)
        reg["Erro"] = linha["erro"]
        reg["Execuções"] = linha["tentativas"]
        reg["Atualizado em"] = linha["atualizado_em"]
        registros.append(reg)
    colunas = ["Processo", "Status"] + [t for _, t in CAMPOS] + ["Erro", "Execuções", "Atualizado em"]
    return pd.DataFrame(registros, columns=colunas)


def _formatar(planilha) -> None:
    cabecalho = PatternFill("solid", fgColor="1F3A5F")
    for celula in planilha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = cabecalho
    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = planilha.dimensions
    for i, coluna in enumerate(planilha.columns, start=1):
        cabecalho_txt, *valores = coluna
        # +4 no cabeçalho para a seta do filtro não cobrir o texto
        largura_cab = len(str(cabecalho_txt.value or "")) + 4
        largura_dados = max((len(str(c.value)) + 2 for c in valores if c.value is not None), default=0)
        planilha.column_dimensions[get_column_letter(i)].width = min(max(largura_cab, largura_dados, 10), 60)


def gerar(linhas: list[dict], caminho: Path) -> Path:
    df = _tabela(linhas)
    concluidos = df[df["Status"].isin(CONCLUIDOS)].drop(columns=["Erro"])
    falhas = df[~df["Status"].isin(CONCLUIDOS)][["Processo", "Status", "Erro", "Execuções", "Atualizado em"]]
    resumo = df["Status"].value_counts().rename_axis("Status").reset_index(name="Quantidade")
    resumo.loc[len(resumo)] = ["total", int(resumo["Quantidade"].sum())]

    caminho.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(caminho, engine="openpyxl") as escritor:
        resumo.to_excel(escritor, sheet_name="Resumo", index=False)
        concluidos.to_excel(escritor, sheet_name="Processos", index=False)
        falhas.to_excel(escritor, sheet_name="Falhas", index=False)
        for planilha in escritor.book.worksheets:
            _formatar(planilha)
    return caminho
