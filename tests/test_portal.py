"""Testes ponta a ponta: robô Playwright contra o portal simulado, rodando local."""

import openpyxl
import pytest

from portal_simulado.servidor import iniciar, processo_existe
from robo import cnj
from robo.main import main


@pytest.fixture
def portal():
    # Falha de 30% e sessão curta para forçar retry e relogin dentro do teste.
    servidor = iniciar(0, taxa_falha=0.3, atraso_max=0.2, consultas_por_sessao=3, semente=42)
    yield f"http://127.0.0.1:{servidor.server_address[1]}"
    servidor.shutdown()


def numeros(qtd):
    return [cnj.gerar(1000 + i * 37, 2024) for i in range(qtd)]


def rodar(portal, tmp_path, entrada, *extra):
    return main(["--fonte", "portal", "--portal-url", portal, "--entrada", str(entrada),
                 "--saida", str(tmp_path / "saida"), *extra])


def test_lote_completo_com_falhas_do_portal(portal, tmp_path):
    lista = numeros(8)
    entrada = tmp_path / "processos.txt"
    entrada.write_text("\n".join([n.formatado for n in lista] + ["0001234-99.2024.8.21.0001"]), encoding="utf-8")

    codigo = rodar(portal, tmp_path, entrada)
    assert codigo == 0  # nenhuma falha definitiva apesar dos 503 e da sessão expirando

    relatorios = list((tmp_path / "saida").glob("relatorio_portal_*.xlsx"))
    assert len(relatorios) == 1
    livro = openpyxl.load_workbook(relatorios[0])
    processos = {linha[0]: linha[1] for linha in livro["Processos"].iter_rows(min_row=2, values_only=True)}
    for n in lista:
        assert processos[n.formatado] == ("ok" if processo_existe(n.digitos) else "nao_encontrado")
    falhas = [linha for linha in livro["Falhas"].iter_rows(min_row=2, values_only=True)]
    assert falhas[0][1] == "numero_invalido"

    pdfs = {p.stem for p in (tmp_path / "saida" / "pdfs").glob("*.pdf")}
    assert pdfs == {n.digitos for n in lista if processo_existe(n.digitos)}
    for p in (tmp_path / "saida" / "pdfs").glob("*.pdf"):
        assert p.read_bytes().startswith(b"%PDF")


def test_segunda_execucao_nao_repete_processos(portal, tmp_path):
    entrada = tmp_path / "processos.txt"
    entrada.write_text("\n".join(n.formatado for n in numeros(4)), encoding="utf-8")
    rodar(portal, tmp_path, entrada)

    rodar(portal, tmp_path, entrada)
    log = (tmp_path / "saida" / "robo.log").read_text(encoding="utf-8")
    assert log.count("já concluído em execução anterior") == 4
    assert "pendentes=0" in log


def test_senha_errada_aborta_o_lote(portal, tmp_path, monkeypatch):
    monkeypatch.setenv("PORTAL_SENHA", "errada")
    entrada = tmp_path / "processos.txt"
    entrada.write_text("\n".join(n.formatado for n in numeros(5)), encoding="utf-8")
    assert rodar(portal, tmp_path, entrada, "--sem-pdf") == 2
    log = (tmp_path / "saida" / "robo.log").read_text(encoding="utf-8")
    assert log.count("fazendo login") == 1  # não ficou tentando a senha errada
